from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class ExtractedSection:
    locator: str
    text: str


def _clean(text: str) -> str:
    return " ".join(text.split()).strip()


def extract_pdf(path: Path) -> list[ExtractedSection]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    sections = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            sections.append(ExtractedSection(f"page {page_number}", text))
    return sections


def extract_pptx(path: Path) -> list[ExtractedSection]:
    from pptx import Presentation

    presentation = Presentation(path)
    sections = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        text = _clean("\n".join(parts))
        if text:
            sections.append(ExtractedSection(f"slide {slide_number}", text))
    return sections


def extract_docx(path: Path) -> list[ExtractedSection]:
    from docx import Document

    document = Document(path)
    sections: list[ExtractedSection] = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean(paragraph.text)
        if text:
            sections.append(ExtractedSection(f"paragraph {number}", text))
    for table_number, table in enumerate(document.tables, start=1):
        rows = [
            " | ".join(_clean(cell.text) for cell in row.cells) for row in table.rows
        ]
        text = _clean("\n".join(rows))
        if text:
            sections.append(ExtractedSection(f"table {table_number}", text))
    return sections


def extract_text(path: Path) -> list[ExtractedSection]:
    text = path.read_text(encoding="utf-8", errors="replace")
    sections: list[ExtractedSection] = []
    buffer: list[str] = []
    section_number = 1
    for line in text.splitlines():
        if line.lstrip().startswith("#") and buffer:
            cleaned = _clean("\n".join(buffer))
            if cleaned:
                sections.append(ExtractedSection(f"section {section_number}", cleaned))
                section_number += 1
            buffer = []
        buffer.append(line)
    cleaned = _clean("\n".join(buffer))
    if cleaned:
        sections.append(ExtractedSection(f"section {section_number}", cleaned))
    return sections


def extract_file(path: Path) -> list[ExtractedSection]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".txt", ".md"}:
        return extract_text(path)
    raise ValueError(f"Unsupported course file: {path.name}")


def iter_course_files(course_dir: Path):
    for path in sorted(course_dir.rglob("*")):
        if not path.is_file() or path.name.startswith((".", "~$")):
            continue
        if path.parent == course_dir and path.name.lower() == "readme.md":
            continue
        if path.is_symlink():
            continue
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path
