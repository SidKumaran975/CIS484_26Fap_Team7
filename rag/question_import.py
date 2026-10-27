"""Parse faculty-authored oral-exam questions from DOCX, XLSX, and CSV files."""

from __future__ import annotations

import csv
import io
import re
import uuid
from pathlib import Path

from docx import Document
from docx.table import Table

QUESTION_PREFIX = re.compile(
    r"^(?:question\s*\d+|q\s*\d+|\d+)\s*[.):\-]\s*",
    re.IGNORECASE,
)
KEY_POINTS_PREFIX = re.compile(
    r"^(?:key\s*points?|answer|expected\s*answer|rubric)\s*:\s*",
    re.IGNORECASE,
)
FOLLOW_UP_PREFIX = re.compile(
    r"^(?:follow[ -]?up|probe)\s*:\s*",
    re.IGNORECASE,
)
BULLET_PREFIX = re.compile(r"^(?:[•\-*]|\d+[.)])\s*")


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split()).strip()


def _split_points(value: str) -> list[str]:
    lines = re.split(r"[\r\n]+|\s*;\s*", value)
    return [
        cleaned
        for line in lines
        if (cleaned := _clean(BULLET_PREFIX.sub("", line)))
    ]


def _question(value: str) -> str:
    return _clean(QUESTION_PREFIX.sub("", value))


def _new_question(
    text: str,
    filename: str,
    key_points=None,
    follow_up="",
    source_kind: str = "Word document",
) -> dict:
    item = {
        "id": "q_" + uuid.uuid4().hex[:8],
        "question": _question(text),
        "key_points": list(key_points or []),
        "source_file": filename,
        "source_excerpt": f"Imported from a faculty-authored {source_kind}.",
    }
    if follow_up:
        item["follow_up"] = _clean(follow_up)
    return item


def _table_questions(table: Table, filename: str) -> list[dict]:
    rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []

    headers = [cell.lower() for cell in rows[0]]
    question_column = next(
        (i for i, value in enumerate(headers) if "question" in value), None
    )
    key_column = next(
        (
            i
            for i, value in enumerate(headers)
            if any(word in value for word in ("key point", "answer", "rubric"))
        ),
        None,
    )
    follow_column = next(
        (i for i, value in enumerate(headers) if "follow" in value or "probe" in value),
        None,
    )
    data_rows = rows[1:] if question_column is not None else rows
    question_column = 0 if question_column is None else question_column

    output = []
    for row in data_rows:
        if question_column >= len(row):
            continue
        text = _question(row[question_column])
        if not text:
            continue
        points = _split_points(row[key_column]) if key_column is not None and key_column < len(row) else []
        follow_up = row[follow_column] if follow_column is not None and follow_column < len(row) else ""
        output.append(_new_question(text, filename, points, follow_up))
    return output


def _finalize_questions(questions: list[dict], limit: int) -> list[dict]:
    unique: list[dict] = []
    seen: set[str] = set()
    for item in questions:
        normalized = item["question"].casefold()
        if len(item["question"]) < 5 or normalized in seen:
            continue
        seen.add(normalized)
        item["key_points"] = list(dict.fromkeys(item["key_points"]))
        unique.append(item)

    if not unique:
        raise ValueError("No questions were found in the uploaded file.")
    if len(unique) > limit:
        raise ValueError(f"The file contains {len(unique)} questions; the limit is {limit}.")
    return unique


def parse_question_docx(data: bytes, filename: str, limit: int = 20) -> list[dict]:
    """Return editable exam-question dictionaries from a DOCX byte payload.

    Supported layouts:
    - one question per paragraph, optionally numbered or labelled Q1/Question 1;
    - ``Key points:`` and ``Follow-up:`` paragraphs after a question;
    - tables with Question, Key Points/Answer/Rubric, and Follow-up columns.
    """
    if Path(filename).suffix.lower() != ".docx":
        raise ValueError("Premade questions must be uploaded as a .docx file.")
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx exposes several malformed-package errors
        raise ValueError("The Word document could not be opened. Save it again as DOCX.") from exc

    questions: list[dict] = []
    current: dict | None = None
    collecting_points = False

    for block in document.iter_inner_content():
        if isinstance(block, Table):
            questions.extend(_table_questions(block, filename))
            current = None
            collecting_points = False
            continue

        text = _clean(block.text)
        if not text:
            continue
        if current is not None and KEY_POINTS_PREFIX.match(text):
            current["key_points"].extend(
                _split_points(KEY_POINTS_PREFIX.sub("", text))
            )
            collecting_points = True
            continue
        if current is not None and FOLLOW_UP_PREFIX.match(text):
            current["follow_up"] = _clean(FOLLOW_UP_PREFIX.sub("", text))
            collecting_points = False
            continue
        if current is not None and collecting_points and BULLET_PREFIX.match(text):
            current["key_points"].extend(_split_points(text))
            continue
        if text.endswith("?") or QUESTION_PREFIX.match(text):
            current = _new_question(text, filename)
            questions.append(current)
            collecting_points = False

    try:
        return _finalize_questions(questions, limit)
    except ValueError as exc:
        if not questions:
            raise ValueError(
                "No questions were found. Put each question in its own paragraph, "
                "number it, or use a table with a Question column."
            ) from exc
        raise


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


HEADER_ALIASES = {
    "question": {"question", "question text", "prompt"},
    "key_points": {
        "key point",
        "key points",
        "answer",
        "expected answer",
        "rubric",
    },
    "follow_up": {"follow up", "follow up question", "probe"},
}


def _cell_text(value) -> str:
    if value is None:
        return ""
    # Preserve line breaks until key-point splitting; questions/follow-ups are
    # normalized by _new_question after their raw cells have been validated.
    return str(value).replace("\xa0", " ").strip()


def _tabular_questions(
    rows,
    filename: str,
    source_kind: str,
    limit: int,
) -> list[dict]:
    iterator = iter(rows)
    header = None
    for raw_row in iterator:
        candidate = [_cell_text(value) for value in raw_row]
        if any(candidate):
            header = candidate
            break
    if header is None:
        raise ValueError("The uploaded file is empty.")

    columns: dict[str, int] = {}
    for index, value in enumerate(header):
        normalized = _normalize_header(value)
        for field, aliases in HEADER_ALIASES.items():
            if normalized in aliases and field not in columns:
                columns[field] = index
    if "question" not in columns:
        raise ValueError(
            "The first nonempty row must contain a Question column. Optional "
            "columns are Key Points and Follow-up."
        )

    questions: list[dict] = []
    for row_number, raw_row in enumerate(iterator, start=2):
        row = list(raw_row)

        def value(field: str) -> str:
            index = columns.get(field)
            return _cell_text(row[index]) if index is not None and index < len(row) else ""

        question = value("question")
        key_points = value("key_points")
        follow_up = value("follow_up")
        if not question and not key_points and not follow_up:
            continue
        if any(text.startswith("=") for text in (question, key_points, follow_up)):
            raise ValueError(
                f"Row {row_number} contains a formula. Convert question fields to plain text."
            )
        if not question:
            raise ValueError(f"Row {row_number} has supporting information but no question.")
        if len(question) > 2000 or len(key_points) > 10000 or len(follow_up) > 2000:
            raise ValueError(f"Row {row_number} contains text that is too long.")
        questions.append(
            _new_question(
                question,
                filename,
                _split_points(key_points),
                follow_up,
                source_kind,
            )
        )
        if len(questions) > limit:
            raise ValueError(f"The file contains more than {limit} questions.")
    return _finalize_questions(questions, limit)


def parse_question_csv(data: bytes, filename: str, limit: int = 20) -> list[dict]:
    if Path(filename).suffix.lower() != ".csv":
        raise ValueError("CSV question lists must use the .csv extension.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Save the CSV file with UTF-8 encoding and upload it again.") from exc
    try:
        rows = csv.reader(io.StringIO(text, newline=""))
        return _tabular_questions(rows, filename, "CSV file", limit)
    except csv.Error as exc:
        raise ValueError("The CSV file could not be parsed.") from exc


def parse_question_xlsx(data: bytes, filename: str, limit: int = 20) -> list[dict]:
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("Excel question lists must use the .xlsx extension.")
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise ValueError(
            "The Excel workbook could not be opened. Save it again as XLSX."
        ) from exc
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not visible:
            raise ValueError("The Excel workbook has no visible worksheets.")
        worksheet = next(
            (sheet for sheet in visible if sheet.title.casefold() == "questions"),
            visible[0],
        )
        rows = worksheet.iter_rows(
            min_row=1, max_row=100, min_col=1, max_col=20, values_only=True
        )
        return _tabular_questions(rows, filename, "Excel workbook", limit)
    finally:
        workbook.close()


def parse_question_file(data: bytes, filename: str, limit: int = 20) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension == ".docx":
        return parse_question_docx(data, filename, limit)
    if extension == ".xlsx":
        return parse_question_xlsx(data, filename, limit)
    if extension == ".csv":
        return parse_question_csv(data, filename, limit)
    raise ValueError("Premade questions must be uploaded as DOCX, XLSX, or CSV.")
