"""Isolated API/routing tests: no real accounts, models, or database touched."""

import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest
import wave
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from PIL import Image
from docx import Document

from rag.config import Settings
from rag.media import MediaGenerator
from rag.profiles import profile_paths

ROOT = Path(__file__).resolve().parent.parent


def photo(color="red"):
    output = io.BytesIO()
    Image.new("RGB", (256, 256), color).save(output, format="PNG")
    return output.getvalue()


def voice():
    output = io.BytesIO()
    with wave.open(output, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(16000)
        audio.writeframes(b"\x10\x00" * 16000 * 6)
    return output.getvalue()


def question_document():
    output = io.BytesIO()
    document = Document()
    document.add_paragraph("Question 1: What is atomicity?")
    document.add_paragraph("Key points: all-or-nothing; rollback on failure")
    document.save(output)
    return output.getvalue()


class ProfessorMediaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        cls.environment = patch.dict(os.environ, {
            "APP_DB_PATH": str(cls.root / "faculty.db"),
            "DB_PATH": str(cls.root / "faculty.db"),
            "FACULTY_EXAMS_DIR": str(cls.root / "exams"),
            "FACULTY_COURSE_STORAGE": str(cls.root / "courses"),
            "FACULTY_INDEX_STORAGE": str(cls.root / "indexes"),
            "FACULTY_MEDIA_STORAGE": str(cls.root / "profiles"),
            "ANALYTICS_DB_PATH": str(cls.root / "analytics.db"),
            "SESSION_DB_PATH": str(cls.root / "sessions.db"),
        })
        cls.environment.start()
        sys.path.insert(0, str(ROOT / "faculty-portal"))
        spec = importlib.util.spec_from_file_location("profile_test_portal", ROOT / "faculty-portal/app.py")
        cls.portal = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.portal)
        for name, role in (("prof_a", "faculty"), ("prof_b", "faculty"), ("student", "student")):
            cls.portal.auth.create_user(name, name, role, "test-password-123")

    @classmethod
    def tearDownClass(cls):
        cls.portal.ANALYTICS.close()
        cls.environment.stop()
        sys.path.remove(str(ROOT / "faculty-portal"))
        cls.temp.cleanup()

    def client(self, name):
        client = TestClient(self.portal.app)
        response = client.post("/api/login", json={"username": name, "password": "test-password-123"})
        self.assertEqual(response.status_code, 200)
        self.addCleanup(client.close)
        return client

    def upload(self, client, color="red", audio=None):
        return client.post("/api/profile/media", files={
            "photo": ("face.png", photo(color), "image/png"),
            "voice": ("voice.wav", voice() if audio is None else audio, "audio/wav"),
        })

    def seed_exam(self, client, exam_id):
        course = client.post("/api/courses", json={"code": "TEST", "name": "Test course"}).json()
        cid = course["id"]
        manifest = self.portal.course_index_dir(cid).parent / "manifest.json"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps({"course_id": cid, "embedding_model": self.portal.EMBED_MODEL}))
        with self.portal.db() as db:
            db.execute("INSERT INTO exams (id,course_id,status) VALUES (?,?,'draft')", (exam_id, cid))
        return {"course_id": cid, "title": "Test", "num_questions": 1,
                "questions": [{"id": "q1", "question": "Explain TCP", "key_points": ["Reliable"]}]}

    def test_access_validation_and_private_previews(self):
        self.assertEqual(TestClient(self.portal.app).get("/api/profile/media").status_code, 401)
        student = self.client("student")
        self.assertEqual(self.upload(student).status_code, 403)
        a, b = self.client("prof_a"), self.client("prof_b")
        self.assertEqual(self.upload(a, audio=b"not a WAV").status_code, 400)
        oversized = a.post("/api/profile/media", files={"photo": ("x.png", b"x" * (10 * 1024 * 1024 + 1)), "voice": ("x.wav", voice())})
        self.assertEqual(oversized.status_code, 413)
        self.assertEqual(self.upload(a).status_code, 200)
        self.assertEqual(self.upload(b, "blue").status_code, 200)
        self.assertNotEqual(a.get("/api/profile/media/photo").content, b.get("/api/profile/media/photo").content)
        self.assertEqual(student.get("/api/profile/media/voice").status_code, 403)
        self.assertEqual(a.get("/api/profile/media/voice").content, voice())

    def test_faculty_imports_premade_questions_as_draft(self):
        faculty = self.client("prof_a")
        course = faculty.post(
            "/api/courses", json={"code": "DOCX", "name": "Import Test"}
        ).json()
        response = faculty.post(
            "/api/exams/import-docx",
            data={
                "course_id": course["id"],
                "title": "Premade Midterm",
                "seconds_per_question": "180",
            },
            files={
                "file": (
                    "questions.docx",
                    question_document(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        imported = response.json()
        self.assertEqual(imported["question_count"], 1)
        self.assertEqual(imported["questions"][0]["key_points"], ["all-or-nothing", "rollback on failure"])
        draft = faculty.get(f"/api/exams/{imported['exam_id']}/edit").json()
        self.assertEqual(draft["status"], "draft")
        self.assertEqual(draft["questions"][0]["question"], "What is atomicity?")

        csv_response = faculty.post(
            "/api/exams/import",
            data={"course_id": course["id"], "title": "CSV Quiz"},
            files={
                "file": (
                    "questions.csv",
                    b"Question,Key Points\nHow does HTTPS work?,TLS; certificates\n",
                    "text/csv",
                )
            },
        )
        self.assertEqual(csv_response.status_code, 200, csv_response.text)
        self.assertEqual(csv_response.json()["question_count"], 1)

    def test_publish_snapshots_replacement_and_ownership(self):
        a, b = self.client("prof_a"), self.client("prof_b")
        first = self.upload(a).json()["version"]
        eid = "exam_0123456789"
        payload = self.seed_exam(a, eid)
        self.assertEqual(b.post(f"/api/exams/{eid}/publish", json=payload).status_code, 404)
        self.assertEqual(a.post(f"/api/exams/{eid}/publish", json=payload).status_code, 200)
        exam_path = self.portal.EXAMS_DIR / f"{eid}.json"
        saved = json.loads(exam_path.read_text())
        self.assertEqual(saved["examiner"]["media_version"], first)
        second = self.upload(a, "green").json()["version"]
        self.assertNotEqual(first, second)
        self.assertEqual(a.get(f'/api/exams/{eid}').json()['examiner']['media_version'],second)
        with self.portal.db() as db:
            owner = db.execute("SELECT id FROM users WHERE username='prof_a'").fetchone()[0]
            self.assertEqual(db.execute('SELECT count(*) FROM AVATAR WHERE UserID=?',(owner,)).fetchone()[0],1)
        self.assertEqual(json.loads(exam_path.read_text()), saved)
        self.assertTrue(all(p.is_file() for p in profile_paths(saved["professor_id"], first)))
        with self.portal.db() as db:
            row = db.execute("SELECT professor_id,media_version FROM exams WHERE id=?", (eid,)).fetchone()
        self.assertEqual(row["media_version"], first)
        self.assertEqual(a.post(f"/api/exams/{eid}/publish", json=payload).status_code, 200)
        self.assertEqual(json.loads(exam_path.read_text())["examiner"]["media_version"], second)
        self.portal.init_db()  # Migration is safe to repeat.

    def test_user_identity_fields_are_saved(self):
        self.portal.auth.create_user('identity-test','', 'student','test-password-123',
                                     first_name='Taylor',last_name='Jones',email='taylor@example.test')
        with self.portal.db() as db:
            row = db.execute('SELECT FirstName,LastName,Email,DisplayName FROM "USER" WHERE UserName=?',('identity-test',)).fetchone()
        self.assertEqual(tuple(row),('Taylor','Jones','taylor@example.test','Taylor Jones'))

    def test_user_full_name_is_split_into_first_and_last(self):
        for index,(display,first,last) in enumerate((
            ('John Smith','',''),('  John   Smith  ','',''),
            ('','John Smith',''),('John Smith','John',''),
        )):
            username=f'full-name-{index}'
            self.portal.auth.create_user(username,display,'student','test-password-123',first_name=first,last_name=last)
            with self.portal.db() as c:
                row=c.execute('SELECT FirstName,LastName FROM "USER" WHERE UserName=?',(username,)).fetchone()
            self.assertEqual(tuple(row),('John','Smith'))

    def test_saved_student_audio_can_be_played_only_by_owning_faculty(self):
        faculty=self.client('prof_a')
        eid='exam_5555555555'
        payload=self.seed_exam(faculty,eid)
        store=self.portal.TRANSCRIPTS
        session=store.create(payload['course_id'],eid,initial_prompt='Explain TCP')
        recording=self.root/'student-answer.webm'
        recording.write_bytes(b'recorded student answer')
        turn=store.record_turn_and_update(session['id'],'q1','Explain TCP','Reliable delivery',{},[],
                                         'move_on',1,0,True,recording={'path':str(recording),'format':'audio/webm'})
        report=faculty.get('/api/transcripts').json()
        saved=next(s for s in report['sessions'] if s['id']==session['id'])
        audio=saved['messages'][1]['recording']
        self.assertEqual(audio['url'],f'/api/recordings/{turn}')
        response=faculty.get(audio['url'])
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.content,recording.read_bytes())
        self.assertEqual(self.client('prof_b').get(audio['url']).status_code,404)
        self.assertEqual(self.client('student').get(audio['url']).status_code,403)
        with TestClient(self.portal.app) as anonymous:
            self.assertEqual(anonymous.get(audio['url']).status_code,401)

    def test_student_upload_is_retained_and_linked_after_transcription(self):
        faculty=self.client('prof_a')
        eid='exam_6666666666'
        payload=self.seed_exam(faculty,eid)
        store=self.portal.TRANSCRIPTS
        session=store.create(payload['course_id'],eid,initial_prompt='Explain TCP')
        def submit(session_id,transcript,recording):
            turn=store.record_turn_and_update(session_id,'q1','Explain TCP',transcript,{},[],
                                             'move_on',1,0,True,recording=recording)
            return {'turn_id':turn,'question_id':'q1','completed':True}
        engine=SimpleNamespace(submit_response=submit,exam_id=eid,
            assessment=SimpleNamespace(course_id=payload['course_id']),examiner=None,close=lambda:None)
        directory=self.root/'answer-recordings'
        with patch.dict(os.environ,{'RECORDINGS_DIR':str(directory),'RETAIN_AUDIO':'false'}), \
             patch.dict(sys.modules,{'faster_whisper':SimpleNamespace(WhisperModel=object)}):
            spec=importlib.util.spec_from_file_location('test_student_audio_app',ROOT/'speech-to-text/app.py')
            student=importlib.util.module_from_spec(spec)
            spec.loader.exec_module(student)
        student.app.dependency_overrides[student.require_portal_student]=lambda:{'id':123,'role':'student'}
        student.app.state.settings=Settings()
        student.app.state.analytics=self.portal.ANALYTICS
        student.app.state.manager=SimpleNamespace(engine_for_session=lambda *args:engine)
        student.app.state.whisper=object()
        student.app.state.media=None
        # No context manager: do not start real models in this routing test.
        client=TestClient(student.app)
        self.addCleanup(client.close)
        with patch.object(student,'transcribe_audio',return_value='Reliable delivery'):
            response=client.post('/api/answers',data={'session_id':session['id']},
                                 files={'audio':('answer.webm',b'recorded speech','audio/webm')})
        self.assertEqual(response.status_code,200,response.text)
        with self.portal.db() as c:
            row=c.execute('SELECT AudioFilePath,FileFormat FROM RESP_REC WHERE ResponseID=?',
                          (response.json()['turn_id'],)).fetchone()
        self.assertEqual(Path(row['AudioFilePath']).read_bytes(),b'recorded speech')
        self.assertEqual(row['FileFormat'],'audio/webm')

    def test_database_publication_survives_export_loss_and_draft_edits(self):
        faculty = self.client('prof_a')
        eid = 'exam_4444444444'
        payload = self.seed_exam(faculty,eid)
        self.assertEqual(faculty.post(f'/api/exams/{eid}/publish',json=payload).status_code,200)
        published = faculty.get(f'/api/exams/{eid}').json()
        (self.portal.EXAMS_DIR/f'{eid}.json').unlink()
        with self.portal.db() as c:
            c.execute('UPDATE exams SET questions=? WHERE id=?', (json.dumps([{'id':'new','question':'draft change'}]),eid))
        self.assertEqual(faculty.get(f'/api/exams/{eid}').json(),published)
        student = self.client('student')
        response = student.get(f'/api/exams/{eid}')
        self.assertEqual(response.status_code,200)
        self.assertNotIn('key_points',response.json()['questions'][0])

    def test_legacy_publish_and_media_routing_cache(self):
        # A newly created faculty account without uploads still uses legacy media.
        self.portal.auth.create_user("legacy", "Legacy", "faculty", "test-password-123")
        legacy = self.client("legacy")
        payload = self.seed_exam(legacy, "exam_1111111111")
        self.assertEqual(legacy.post("/api/exams/exam_1111111111/publish", json=payload).status_code, 200)
        self.assertIsNone(json.loads((self.portal.EXAMS_DIR / "exam_1111111111.json").read_text())["examiner"])
        a, b = self.client("prof_a"), self.client("prof_b")
        va, vb = self.upload(a).json()["version"], self.upload(b, "blue").json()["version"]
        with self.portal.db() as db:
            ids = dict(db.execute("SELECT username,id FROM users"))
        generator = MediaGenerator(Settings(media_dir=self.root / "generated", media_mode="avatar"))
        commands = []
        def render(command, cwd, label):
            commands.append(command)
            output_flag = "--output" if "--output" in command else "--outfile"
            Path(command[command.index(output_flag) + 1]).write_bytes(b"media")
        ea = {"professor_id": ids["prof_a"], "media_version": va}
        eb = {"professor_id": ids["prof_b"], "media_version": vb}
        with patch.object(generator, "_run", side_effect=render):
            url_a = generator.generate("Question", ea)
            self.assertEqual(generator.generate("Question", ea), url_a)
            url_b = generator.generate("Question", eb)
            legacy_url = generator.generate("Question")
        self.assertEqual(len(commands), 6)
        self.assertEqual(len({url_a, url_b, legacy_url}), 3)
        face, audio = profile_paths(ids["prof_a"], va)
        self.assertEqual(commands[0][commands[0].index("--reference-audio") + 1], str(audio))
        self.assertEqual(commands[1][commands[1].index("--face") + 1], str(face))
        with self.assertRaises(ValueError):
            generator.generate("Question", {"professor_id": ids["prof_a"], "media_version": "0" * 64})
        with self.assertRaises(ValueError):
            profile_paths(1, "../../private")

    def test_faculty_can_only_read_transcripts_for_owned_exams(self):
        a, b = self.client("prof_a"), self.client("prof_b")
        exam_a, exam_b = "exam_2222222222", "exam_3333333333"
        self.seed_exam(a, exam_a)
        self.seed_exam(b, exam_b)
        session_a = self.portal.TRANSCRIPTS.create(
            "course-a",
            exam_a,
            student_id=3,
            student_username="student",
            student_name="Test Student",
            initial_prompt="Explain TCP.",
        )
        self.portal.TRANSCRIPTS.record_turn_and_update(
            session_id=session_a["id"], question_id="q1",
            prompt="Explain TCP.", transcript="TCP provides reliable delivery.",
            evaluation={}, retrieved=[], action="move_on", question_index=1,
            follow_up_count=0, completed=True,
        )
        self.portal.TRANSCRIPTS.create(
            "course-b", exam_b, student_name="Other Student",
            initial_prompt="Private prompt",
        )

        response = a.get("/api/transcripts")
        self.assertEqual(response.status_code, 200)
        sessions = response.json()["sessions"]
        self.assertEqual(len(sessions), 1)
        self.assertEqual(sessions[0]["exam_id"], exam_a)
        self.assertEqual(sessions[0]["student_name"], "Test Student")
        self.assertEqual(
            sessions[0]["messages"][1]["text"],
            "TCP provides reliable delivery.",
        )
        self.assertEqual(self.client("student").get("/api/transcripts").status_code, 403)


if __name__ == "__main__":
    unittest.main()
