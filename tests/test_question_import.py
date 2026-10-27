import io
import unittest

from docx import Document
from openpyxl import Workbook

from rag.question_import import (
    parse_question_csv,
    parse_question_docx,
    parse_question_file,
    parse_question_xlsx,
)


def document_bytes(build):
    document = Document()
    build(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def workbook_bytes(build):
    workbook = Workbook()
    build(workbook)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class QuestionImportTests(unittest.TestCase):
    def test_sprint_one_paragraph_format(self):
        data = document_bytes(
            lambda document: [
                document.add_paragraph(text)
                for text in (
                    "What is ACID compliance?",
                    "2. Explain the difference between a process and a thread.",
                    "Not a question or numbered prompt",
                )
            ]
        )
        questions = parse_question_docx(data, "Question List.docx")
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[1]["question"], "Explain the difference between a process and a thread.")

    def test_key_points_follow_up_and_table(self):
        def build(document):
            document.add_paragraph("Question 1: What does atomicity mean?")
            document.add_paragraph("Key points: all-or-nothing; rollback on failure")
            document.add_paragraph("Follow-up: What happens after one operation fails?")
            table = document.add_table(rows=2, cols=3)
            table.rows[0].cells[0].text = "Question"
            table.rows[0].cells[1].text = "Key Points"
            table.rows[0].cells[2].text = "Follow-up"
            table.rows[1].cells[0].text = "How does HTTPS protect traffic?"
            table.rows[1].cells[1].text = "TLS encryption; certificates"
            table.rows[1].cells[2].text = "What does the certificate prove?"

        questions = parse_question_docx(document_bytes(build), "exam.docx")
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]["key_points"], ["all-or-nothing", "rollback on failure"])
        self.assertEqual(questions[1]["follow_up"], "What does the certificate prove?")

    def test_rejects_wrong_extension_and_empty_document(self):
        data = document_bytes(lambda document: document.add_paragraph("Course notes only."))
        with self.assertRaisesRegex(ValueError, "DOCX|docx"):
            parse_question_docx(data, "questions.doc")
        with self.assertRaisesRegex(ValueError, "No questions"):
            parse_question_docx(data, "questions.docx")

    def test_csv_columns_and_multiline_key_points(self):
        data = (
            "Question,Key Points,Follow-up\n"
            '"What is atomicity?","all-or-nothing\nrollback on failure","What happens after failure?"\n'
            '"How does HTTPS work?","TLS; certificates",""\n'
        ).encode()
        questions = parse_question_csv(data, "questions.csv")
        self.assertEqual(len(questions), 2)
        self.assertEqual(
            questions[0]["key_points"], ["all-or-nothing", "rollback on failure"]
        )
        self.assertEqual(questions[1]["source_file"], "questions.csv")

    def test_xlsx_prefers_questions_sheet_and_alias_headers(self):
        def build(workbook):
            workbook.active.title = "Notes"
            workbook.active.append(["Not", "Questions"])
            sheet = workbook.create_sheet("Questions")
            sheet.append(["Prompt", "Expected Answer", "Probe"])
            sheet.append(
                [
                    "What is a database index?",
                    "Speeds lookup; trades storage and write cost",
                    "When might an index hurt performance?",
                ]
            )

        questions = parse_question_xlsx(workbook_bytes(build), "questions.xlsx")
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["question"], "What is a database index?")
        self.assertEqual(len(questions[0]["key_points"]), 2)

    def test_tabular_import_rejects_formulas_and_missing_question_header(self):
        formula_csv = b"Question,Key Points\n=1+1,answer\n"
        with self.assertRaisesRegex(ValueError, "formula"):
            parse_question_csv(formula_csv, "questions.csv")
        with self.assertRaisesRegex(ValueError, "Question column"):
            parse_question_csv(b"Topic,Answer\nACID,Atomicity\n", "questions.csv")

    def test_file_dispatch_and_unsupported_format(self):
        questions = parse_question_file(
            b"Question,Key Points\nWhat is ACID?,Atomicity\n", "questions.csv"
        )
        self.assertEqual(len(questions), 1)
        with self.assertRaisesRegex(ValueError, "DOCX, XLSX, or CSV"):
            parse_question_file(b"question", "questions.xls")


if __name__ == "__main__":
    unittest.main()
